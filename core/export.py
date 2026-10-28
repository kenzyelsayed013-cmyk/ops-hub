"""Excel + PDF export with the same colour highlighting as the screen."""
import io
import re
import numpy as np
import pandas as pd
from . import insights as I

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    HAS_PDF = True
except Exception:
    HAS_PDF = False

BG = {"good": "DCFCE7", "warn": "FEF3C7", "bad": "FEE2E2"}
FG = {"good": "14532D", "warn": "78350F", "bad": "7F1D1D"}
CSS_TONE = {"DCFCE7": "good", "FEF3C7": "warn", "FEE2E2": "bad"}
FONT, BOLD, UNICODE = "Helvetica", "Helvetica-Bold", False


def _fonts():
    global FONT, BOLD, UNICODE
    if not HAS_PDF:
        return
    for base in ("/usr/share/fonts/truetype/dejavu/", "/usr/share/fonts/dejavu/", "/usr/share/fonts/TTF/"):
        try:
            pdfmetrics.registerFont(TTFont("DejaVu", base + "DejaVuSans.ttf"))
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", base + "DejaVuSans-Bold.ttf"))
            FONT, BOLD, UNICODE = "DejaVu", "DejaVu-Bold", True
            return
        except Exception:
            continue


_fonts()


def clean(s):
    s = str(s).replace("✅", "Y").replace("❌", "X").replace("·", "-").replace("→", "->").replace("−", "-")
    if not UNICODE:
        s = s.replace("✔", "OK").replace("✖", "X").replace("⚠", "!").replace("★", "*")
        s = s.encode("latin-1", "replace").decode("latin-1")
    return s


def fmt(v, pct=False):
    if isinstance(v, str):
        return clean(v)
    if v is None or (isinstance(v, float) and np.isnan(v)) or (not isinstance(v, (int, float, np.integer, np.floating)) and pd.isna(v)):
        return ""
    if pct:
        return f"{v:.1f}%"
    if isinstance(v, (int, np.integer)):
        return f"{v:,}"
    if isinstance(v, (float, np.floating)):
        return f"{v:,.0f}" if (abs(v) >= 1000 or float(v).is_integer()) else (f"{v:,.2f}" if abs(v) < 100 else f"{v:,.1f}")
    return clean(v)


def tones(blk):
    """{(row, col): 'good'|'warn'|'bad'} for a block."""
    df, out = blk["df"], {}
    rl = {c: r for c, r in {**I.TABLE_RULES, **(blk.get("rules") or {})}.items() if c in df.columns}
    for c, r in rl.items():
        j = list(df.columns).index(c)
        for i, v in enumerate(df[c].tolist()):
            t = I.grade(v, *r)
            if t in BG:
                out[(i, j)] = t
    for c in [c for c in (blk.get("text_cols") or ["Status", "Segment"]) if c in df.columns]:
        j = list(df.columns).index(c)
        for i, v in enumerate(df[c].tolist()):
            t = I.tone_of_text(v)
            if t in BG:
                out[(i, j)] = t
    tn = blk.get("tone")
    if tn is not None:
        arr = tn.values
        for i in range(min(len(df), arr.shape[0])):
            for j in range(min(len(df.columns), arr.shape[1])):
                m = re.search(r"#([0-9A-Fa-f]{6})", str(arr[i, j]))
                if m and m.group(1).upper() in CSS_TONE:
                    out[(i, j)] = CSS_TONE[m.group(1).upper()]
    return out


# ---------------------------------------------------------------- Excel
def xlsx_bytes(blocks):
    from openpyxl.styles import Font, PatternFill
    buf, used = io.BytesIO(), set()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for b in blocks:
            df = b["df"].copy()
            for c in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[c]):
                    df[c] = df[c].dt.date
            nm = re.sub(r"[\[\]:*?/\\]", " ", str(b.get("title") or "data")).strip()[:31] or "data"
            k, base = 2, nm
            while nm.lower() in used:
                nm = f"{base[:28]} {k}"
                k += 1
            used.add(nm.lower())
            df.to_excel(xw, sheet_name=nm, index=False)
            ws = xw.sheets[nm]
            ws.freeze_panes = "A2"
            for i, c in enumerate(df.columns, 1):
                h = ws.cell(row=1, column=i)
                h.font, h.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="14181F")
                width = max([len(str(c))] + [len(str(x)) for x in df[c].head(200).tolist()]) + 2
                ws.column_dimensions[h.column_letter].width = min(40, max(10, width))
            if b.get("first_row") and len(df):
                for i in range(1, len(df.columns) + 1):
                    ws.cell(row=2, column=i).font = Font(bold=True)
            for (r, c), t in tones(b).items():
                cell = ws.cell(row=r + 2, column=c + 1)
                cell.fill, cell.font = PatternFill("solid", fgColor=BG[t]), Font(bold=True, color=FG[t])
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
def pdf_bytes(blocks, title, subtitle=""):
    big = any(len(b["df"].columns) > 14 for b in blocks)
    page = landscape(A3 if big else A4)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20, title=clean(title))
    avail = page[0] - 40
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName=BOLD, fontSize=18, alignment=0)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=BOLD, fontSize=12, backColor=colors.HexColor("#FFC400"), borderPadding=4, spaceBefore=10, spaceAfter=6)
    nm = ParagraphStyle("nm", parent=ss["Normal"], fontName=FONT, fontSize=9, textColor=colors.HexColor("#4B5563"))
    story = [Paragraph(clean(title), h1), Paragraph(clean(subtitle), nm), Spacer(1, 8)]
    for b in blocks:
        df = b["df"]
        if df is None or df.empty:
            continue
        cols = list(df.columns)
        n = len(cols)
        fs = 8 if n <= 9 else (7 if n <= 14 else (6 if n <= 22 else 5))
        hs = ParagraphStyle("hs", fontName=BOLD, fontSize=fs, leading=fs + 2, textColor=colors.white)
        pct = [str(c).endswith("%") for c in cols]
        rows = [[fmt(v, pct[j]) for j, v in enumerate(r)] for r in df.itertuples(index=False, name=None)]
        data = [[Paragraph(clean(c), hs) for c in cols]] + rows
        hw = [max([len(x) for x in clean(c).split()] + [3]) * 1.12 for c in cols]
        lens = [max([hw[j]] + [len(r[j]) for r in rows[:300]]) for j in range(n)]
        w = [l * fs * 0.6 + 8 for l in lens]
        w = [x * (avail / sum(w) if sum(w) > avail else min(avail / sum(w), 1.8)) for x in w]
        t = Table(data, colWidths=w, repeatRows=1, hAlign="LEFT")
        st_ = [("FONTNAME", (0, 0), (-1, -1), FONT), ("FONTSIZE", (0, 0), (-1, -1), fs), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D1D5DB")),
               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14181F")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")])]
        if b.get("first_row"):
            st_ += [("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#14181F")), ("TEXTCOLOR", (0, 1), (-1, 1), colors.white), ("FONTNAME", (0, 1), (-1, 1), BOLD)]
        for (r, c), tn in tones(b).items():
            st_ += [("BACKGROUND", (c, r + 1), (c, r + 1), colors.HexColor("#" + BG[tn])), ("TEXTCOLOR", (c, r + 1), (c, r + 1), colors.HexColor("#" + FG[tn])),
                    ("FONTNAME", (c, r + 1), (c, r + 1), BOLD)]
        t.setStyle(TableStyle(st_))
        if b.get("title"):
            story.append(Paragraph(clean(b["title"]), h2))
        story += [t, Spacer(1, 10)]
    doc.build(story)
    return buf.getvalue()
